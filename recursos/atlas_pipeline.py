"""Atlas electoral: integración reproducible de siete archivos aportados.

Los Excel se leen sin modificarlos. Identidades de personas y partidos son
nominales salvo documentación expresa. Ninguna asociación se interpreta como causal.
"""
from pathlib import Path
from collections import defaultdict, Counter
import hashlib, json, re, unicodedata, math
from datetime import datetime
import numpy as np
import pandas as pd
import openpyxl

FILES = {
 'leg': 'Base_Electoral_Diputaciones_1953-2026(20260918-171611).xlsx',
 'pres': 'Base_Electoral_Presidencial_1953-2026(4).xlsx',
 'cand': 'Base_Final_Candidaturas_Diputaciones_1953-2026_ENTREGA.xlsx',
 'formula': 'Base_Final_Candidaturas_Presidenciales_1953-2026(4).xlsx',
 'electas': 'Base_Final_Diputaciones_Electas_1953-2026(4).xlsx',
 'diseno': 'Base_Final_Diseno_Papeletas_Diputaciones_1953-2026(4).xlsx',
 'cancelados': 'Base_partidos_cancelados.xlsx',
}
YEARS = [1953, 1958] + list(range(1962, 2027, 4))
PROVINCES = ['San José','Alajuela','Cartago','Heredia','Guanacaste','Puntarenas','Limón']

def norm(value):
    if value is None or pd.isna(value): return ''
    return ' '.join(''.join(c for c in unicodedata.normalize('NFD', str(value).upper()) if unicodedata.category(c) != 'Mn').split())

def party(value):
    text = re.sub(r'\([^)]*\)', '', str(value or ''))
    text = re.sub(r'^(PARTIDO|COALICI[ÓO]N)\s+', '', text.strip(), flags=re.I)
    key=norm(text)
    return {'UNIDAD SOCIALCRISTIANA':'UNIDAD SOCIAL CRISTIANA','ORG SOCIALISTA DE LOS TRABAJADORES':'ORGANIZACION SOCIALISTA DE LOS TRABAJADORES'}.get(key,key)

def province(value):
    return {norm(p):p for p in PROVINCES}.get(norm(value), str(value or ''))

def number(value):
    if isinstance(value, (int,float)) and not isinstance(value,bool) and math.isfinite(value): return float(value)
    return None

def ratio(n,d):
    return 100*n/d if n is not None and d is not None and d>0 else None

def records(df):
    return json.loads(df.to_json(orient='records',date_format='iso',force_ascii=False))

def read_sheet(path, name='Datos'):
    w=openpyxl.load_workbook(path,data_only=True,read_only=False)
    rows=list(w[name].values);w.close()
    return pd.DataFrame(rows[1:],columns=rows[0]).dropna(how='all')

def load_sources(root):
    root=Path(root)
    missing=[name for name in FILES.values() if not (root/name).exists()]
    if missing: raise FileNotFoundError('Coloque estas bases junto al notebook o cambie DATA_DIR: '+', '.join(missing))
    manifest=[]
    for role,name in FILES.items():
        b=(root/name).read_bytes()
        manifest.append({'base':role,'archivo':name,'sha256':hashlib.sha256(b).hexdigest(),'bytes':len(b)})
    raw={k:read_sheet(root/v) for k,v in FILES.items() if k not in ('leg','pres')}
    expected=['ID candidatura','Año','Provincia','Partido','Tipo de candidatura','Orden en lista','Nombre completo','Sexo','Electa','Cédula','Alias']
    if list(raw['cand'].columns)!=expected:raise ValueError('La hoja Datos de candidaturas debe contener las 11 columnas exactas descritas en DOCUMENTO_MADRE.md.')
    if raw['cand']['ID candidatura'].duplicated().any():raise ValueError('ID candidatura duplicados en la base de diputaciones.')
    if raw['cand']['Electa'].map(lambda x:norm(x)=='SI').sum()!=len(raw['electas']):
        raise ValueError('Las marcas Electa de candidaturas no concuerdan con el total de diputaciones electas.')
    if len(raw['cand'])!=14874:print('Aviso: base de candidaturas actualizada con',len(raw['cand']),'filas.')
    raw['coincidencias_registrales']=read_sheet(root/FILES['cancelados'],'Coincidencias electorales')
    return raw,manifest

def extract_elections(root):
    votes=[];totals=[];seats=[];issues=[]
    controls={'VOTOS EN BLANCO':'blancos','VOTOS NULOS':'nulos','ABSTENCIONISMO':'abstenciones','TOTAL PADRON':'padron','TOTAL VALIDOS':'validos','TOTAL VOTOS':'emitidos'}
    for role,arena in [('leg','Legislativa'),('pres','Presidencial')]:
        w=openpyxl.load_workbook(Path(root)/FILES[role],data_only=True)
        for s in w:
            if s.title not in map(str,YEARS):continue
            year=int(s.title);rows=list(s.values)
            for i,h in enumerate(rows):
                first=norm(h[0])
                if first.startswith('VOTOS'):
                    rnd=2 if 'SEGUNDA' in first else 1
                    start=next((j for j,x in enumerate(h) if norm(x)=='VOTOS EN BLANCO'),None)
                    if start is None:continue
                    for row in rows[i+1:i+8]:
                        p=province(row[0])
                        if p not in PROVINCES:continue
                        t={'anio':year,'provincia':p,'arena':arena,'ronda':rnd}
                        for j,key in enumerate(h):
                            if norm(key) in controls:t[controls[norm(key)]]=number(row[j])
                        valid_sum=0
                        for j in range(1,start):
                            val=number(row[j])
                            if h[j] is not None and val is not None:
                                votes.append(dict(anio=year,provincia=p,arena=arena,ronda=rnd,partido=party(h[j]),partido_fuente=h[j],votos=val))
                                valid_sum+=val
                        for key in controls.values():t.setdefault(key,None)
                        t['suma_partidos']=valid_sum
                        for typ,delta in [('Votos por partido frente a válidos',valid_sum-(t['validos'] or 0)),('Padrón frente a emitidos y abstenciones',(t['padron'] or 0)-(t['emitidos'] or 0)-(t['abstenciones'] or 0)),('Emitidos frente a válidos, blancos y nulos',(t['emitidos'] or 0)-sum(t[x] or 0 for x in ('validos','blancos','nulos')))]:
                            if abs(delta)>.01:issues.append({'base':role,'anio':year,'provincia':p,'detalle':typ,'diferencia':delta})
                        t['emitidos_fuente']=t['emitidos']
                        if all(t[x] is not None for x in ('validos','blancos','nulos')):
                            t['emitidos']=t['validos']+t['blancos']+t['nulos']
                        totals.append(t)
                elif first=='DIPUTACIONES':
                    end=next((j for j,x in enumerate(h) if norm(x)=='TOTAL'),len(h))
                    for row in rows[i+1:i+8]:
                        p=province(row[0])
                        if p not in PROVINCES:continue
                        for j in range(1,end):
                            v=number(row[j])
                            if h[j] and v is not None:seats.append({'anio':year,'provincia':p,'partido':party(h[j]),'escanos':v})
        w.close()
    v=pd.DataFrame(votes);t=pd.DataFrame(totals);s=pd.DataFrame(seats)
    # Nacional = suma de las siete provincias, nunca promedio de porcentajes.
    keys=['anio','arena','ronda'];cols=list(controls.values())+['suma_partidos','emitidos_fuente']
    nat=t.groupby(keys)[cols].sum(min_count=7).reset_index();nat['provincia']='Nacional'
    t=pd.concat([t,nat],ignore_index=True)
    nv=v.groupby(keys+['partido'],as_index=False).agg(votos=('votos','sum'),partido_fuente=('partido_fuente','first'));nv['provincia']='Nacional'
    v=pd.concat([v,nv],ignore_index=True)
    ns=s.groupby(['anio','partido'],as_index=False).escanos.sum();ns['provincia']='Nacional';s=pd.concat([s,ns],ignore_index=True)
    for col,num,den in [('abstencion_pct','abstenciones','padron'),('participacion_pct','emitidos','padron'),('blancos_pct','blancos','emitidos'),('nulos_pct','nulos','emitidos')]:
        t[col]=100*t[num]/t[den].replace(0,np.nan)
    t['invalidos_pct']=100*(t.blancos+t.nulos)/t.emitidos.replace(0,np.nan)
    v=v.merge(t[keys+['provincia','validos']],on=keys+['provincia'],validate='many_to_one')
    v['porcentaje']=100*v.votos/v.validos.replace(0,np.nan)
    return v,t,s,issues

def prepare_people(raw,issues):
    c=raw['cand'].rename(columns={'ID candidatura':'id','Año':'anio','Provincia':'provincia','Partido':'partido','Tipo de candidatura':'tipo','Orden en lista':'orden','Nombre completo':'nombre','Sexo':'sexo','Electa':'electa','Alias':'alias'})
    e=raw['electas'].rename(columns={'ID registro':'id_registro','Posición provincial':'posicion','Año de elección':'anio','Provincia':'provincia','Partido':'partido','Nombre completo':'nombre','Sexo':'sexo','Período legislativo':'periodo','Primer apellido':'apellido1','Segundo apellido':'apellido2'})
    f=raw['formula'].rename(columns={'Año':'anio','Nombre completo':'nombre','Sexo':'sexo','Partido':'partido','Cargo':'cargo','Electo':'electa','ID fórmula':'formula','Orden en fórmula':'orden','Primer apellido':'apellido1','Segundo apellido':'apellido2'})
    for df in (c,e,f):
        df['partido']=df.partido.map(party);df['nombre_key']=df.nombre.map(norm)
        df['anio']=pd.to_numeric(df.anio).astype(int)
        df['sexo']=df.sexo.fillna('Sin dato')
        if 'provincia' in df:df['provincia']=df.provincia.map(province)
    placeholders=c[c.nombre_key==''].copy()
    if len(placeholders):issues.append({'base':'cand','anio':None,'provincia':'Varias','detalle':'Filas sin nombre excluidas de conteos de personas/candidaturas identificadas','diferencia':len(placeholders)})
    c=c[c.nombre_key!=''].copy();f=f[f.nombre_key!=''].copy()
    c['electa']=c.electa.map(lambda x: norm(x)=='SI');f['electa']=f.electa.map(lambda x:norm(x)=='SI')
    c['cargo']='Diputación';f['tipo']='Fórmula presidencial';f['provincia']='Nacional'
    c['id_evento']=c['id'];f['id_evento']=[f'CR-FORM-{r.anio}-{r.formula}-{r.orden}-{i}' for i,r in enumerate(f.itertuples())]
    c['fuente']='Candidaturas legislativas';f['fuente']='Candidaturas presidenciales'
    # Apellidos solo desde campos expresos; no se trocean nombres históricos.
    surname_map={r.nombre_key:(r.apellido1,r.apellido2) for r in pd.concat([e,f]).itertuples()}
    for df in (c,e,f):
        if 'apellido1' not in df:
            df['apellido1']=df.nombre_key.map(lambda n:surname_map.get(n,(None,None))[0]);df['apellido2']=df.nombre_key.map(lambda n:surname_map.get(n,(None,None))[1])
    known_e={(r.anio,r.provincia,r.partido,r.nombre_key) for r in e.itertuples()}
    c['electa_declaratoria']=[(r.anio,r.provincia,r.partido,r.nombre_key) in known_e for r in c.itertuples()]
    mismatch=c[c.electa!=c.electa_declaratoria]
    for r in mismatch.itertuples():issues.append({'base':'cand/electas','anio':r.anio,'provincia':r.provincia,'detalle':f'Bandera electa y cruce nominal no coinciden: {r.nombre} · {r.partido}','diferencia':None})
    cand_keys={(r.anio,r.provincia,r.partido,r.nombre_key) for r in c.itertuples()}
    missing=e[[ (r.anio,r.provincia,r.partido,r.nombre_key) not in cand_keys for r in e.itertuples() ]]
    # La fuente de electas sigue separada: no inventamos candidaturas faltantes.
    for r in missing.itertuples():issues.append({'base':'electas/cand','anio':r.anio,'provincia':r.provincia,'detalle':'Electa sin coincidencia nominal exacta en candidaturas: '+r.nombre,'diferencia':None})
    c['electa_fuente']=c['electa']
    # El cruce nominal audita: una abreviatura en la lista de electas no
    # convierte en no electa a una candidatura explícitamente marcada Sí.
    events=pd.concat([c,f],ignore_index=True)
    cols=['id_evento','anio','nombre','nombre_key','partido','provincia','cargo','tipo','orden','sexo','electa','fuente','apellido1','apellido2']
    events=events[cols].copy()
    # Un mismo nombre con varias cédulas no se une automáticamente.
    ced=defaultdict(set)
    for r in c.to_dict('records'):
        value=r.get('Cédula')
        if value is not None and pd.notna(value):ced[r['nombre_key']].add(str(value))
    collisions={k for k,v in ced.items() if len(v)>1}
    events['id_persona']=events.nombre_key.map(lambda n:hashlib.sha256(n.encode()).hexdigest()[:14])
    events['identidad']=events.nombre_key.map(lambda n:'Homónimo potencial; no consolidado' if n in collisions else 'Coincidencia nominal exacta')
    for idx in events.index[events.nombre_key.isin(collisions)]:events.at[idx,'id_persona']+='-'+str(idx)
    events=events.sort_values(['anio','nombre','cargo','provincia']).reset_index(drop=True)
    prior={};transitions=[]
    events['experiencia_previa']=False;events['participaciones_previas']=0
    events['cambio_partido']=False;events['cambio_provincia']=False
    for (pid,y),idx in events.groupby(['id_persona','anio'],sort=True).groups.items():
        g=events.loc[idx];prev=prior.get(pid);ps=set(g.partido);provs=set(g.provincia)-{'Nacional'}
        if prev:
            events.loc[idx,'experiencia_previa']=True;events.loc[idx,'participaciones_previas']=prev['n']
            events.loc[idx,'cambio_partido']=not bool(ps & prev['p'])
            events.loc[idx,'cambio_provincia']=bool(provs and prev['prov'] and not provs & prev['prov'])
            for dim,new,old in [('Partido',ps,prev['p']),('Provincia',provs,prev['prov']),('Cargo',set(g.cargo),prev['cargo'])]:
                if len(new)==len(old)==1 and new!=old:transitions.append({'id_persona':pid,'nombre':g.iloc[0]['nombre'],'anio_origen':prev['y'],'anio':y,'dimension':dim,'origen':next(iter(old)),'destino':next(iter(new)),'sexo':g.iloc[0]['sexo']})
        prior[pid]={'n':(prev['n']+1 if prev else 1),'p':ps,'prov':provs or (prev['prov'] if prev else set()),'cargo':set(g.cargo),'y':y}
    people=[]
    for pid,g in events.groupby('id_persona'):
        ys=sorted(g.anio.unique());wins=sorted(g[g.electa].anio.unique());firstwin=wins[0] if wins else None
        people.append({'id_persona':pid,'nombre':g.iloc[-1]['nombre'],'sexo':g.iloc[-1]['sexo'],'inicio':int(ys[0]),'fin':int(ys[-1]),'elecciones':len(ys),'candidaturas':len(g),'partidos':len(set(g.partido)),'provincias':len(set(g.provincia)-{'Nacional'}),'cargos':len(set(g.cargo)),'elecciones_exitosas':len(wins),'primera_eleccion':int(firstwin) if firstwin is not None else None,'intentos_hasta_eleccion':sum(y<=firstwin for y in ys) if firstwin is not None else None,'identidad':g.iloc[0]['identidad']})
    c=c.drop(columns=['Cédula'],errors='ignore')
    return c,e,f,events,pd.DataFrame(people),pd.DataFrame(transitions),placeholders

def calculate_metrics(v,t,s,c,e,events):
    groupcols=['anio','provincia','arena','ronda'];metrics=[]
    for key,g in v.groupby(groupcols):
        vals=g.votos.to_numpy(float);total=vals.sum();shares=np.sort(vals/total)[::-1] if total else np.array([])
        rec=dict(zip(groupcols,key));rec.update({'partidos_con_votos':int((vals>0).sum()),'nep':1/(shares**2).sum() if len(shares) else None,'concentracion_1':100*shares[0] if len(shares) else None,'concentracion_2':100*shares[:2].sum() if len(shares) else None,'margen':100*(shares[0]-(shares[1] if len(shares)>1 else 0)) if len(shares) else None,'ganador':g.loc[g.votos.idxmax(),'partido'] if total else None})
        metrics.append(rec)
    m=pd.DataFrame(metrics).merge(t,on=groupcols,validate='one_to_one')
    # Volatilidad por etiqueta nominal, no identidad jurídica.
    prev={};vol=[]
    for r in m.sort_values('anio').to_dict('records'):
        key=(r['provincia'],r['arena'],r['ronda']);g=v[(v.anio==r['anio'])&(v.provincia==r['provincia'])&(v.arena==r['arena'])&(v.ronda==r['ronda'])]
        shares=dict(zip(g.partido,g.votos/g.votos.sum()*100));old=prev.get(key)
        vol.append((r['anio'],*key, sum(abs(shares.get(p,0)-old.get(p,0)) for p in set(shares)|set(old))/2 if old else None));prev[key]=shares
    vv=pd.DataFrame(vol,columns=['anio','provincia','arena','ronda','volatilidad_nominal']);m=m.merge(vv,on=groupcols)
    base=v[v.ronda==1].pivot_table(index=['anio','provincia','partido'],columns='arena',values=['votos','porcentaje'],aggfunc='sum').reset_index()
    base.columns=['_'.join(x).rstrip('_') if isinstance(x,tuple) else x for x in base.columns]
    for label in ['Legislativa','Presidencial']:
        for f in ['votos','porcentaje']:
            if f+'_'+label not in base:base[f+'_'+label]=np.nan
    base['brecha_pp']=base.porcentaje_Presidencial-base.porcentaje_Legislativa
    base=base.merge(s,on=['anio','provincia','partido'],how='outer',validate='one_to_one')
    # Si la suma de escaños observada coincide con la magnitud se acepta ausencia=0.
    st=s.groupby(['anio','provincia']).escanos.sum().rename('magnitud').reset_index();base=base.merge(st,on=['anio','provincia'],how='left')
    base['escanos']=base.escanos.fillna(0);base['escanos_pct']=100*base.escanos/base.magnitud
    base['desviacion_pp']=base.escanos_pct-base.porcentaje_Legislativa
    candidates=c.copy();owners=candidates[candidates.tipo=='Propietaria']
    stats=[]
    for (y,p,pt),g in candidates.groupby(['anio','provincia','partido']):
        own=g[g.tipo=='Propietaria'];heads=own[own.orden==1]
        stats.append(dict(anio=y,provincia=p,partido=pt,candidaturas=len(g),propietarias=len(own),suplencias=int((g.tipo=='Suplencia').sum()),mujeres=int((g.sexo=='Mujer').sum()),mujeres_propietarias=int((own.sexo=='Mujer').sum()),encabezamientos_mujeres=int((heads.sexo=='Mujer').sum())))
    stats=pd.DataFrame(stats);n=stats.groupby(['anio','partido'],as_index=False).sum(numeric_only=True);n['provincia']='Nacional';stats=pd.concat([stats,n])
    panel=base.merge(stats,on=['anio','provincia','partido'],how='outer',validate='one_to_one')
    panel=panel.merge(st.rename(columns={'magnitud':'magnitud_territorial'}),on=['anio','provincia'],how='left',validate='many_to_one')
    panel['magnitud']=panel.magnitud.fillna(panel.magnitud_territorial)
    panel=panel.drop(columns='magnitud_territorial')
    panel['mujeres_pct']=100*panel.mujeres_propietarias/panel.propietarias.replace(0,np.nan)
    panel['cobertura_lista_pct']=100*panel.propietarias/panel.magnitud
    elected=e.groupby(['anio','provincia','partido']).agg(electas_registro=('nombre','size'),mujeres_electas=('sexo',lambda x:(x=='Mujer').sum())).reset_index()
    ne=elected.groupby(['anio','partido'],as_index=False).sum(numeric_only=True);ne['provincia']='Nacional';elected=pd.concat([elected,ne]);panel=panel.merge(elected,on=['anio','provincia','partido'],how='outer',validate='one_to_one')
    for (y,p),g in panel.groupby(['anio','provincia']):
        mask=(m.anio==y)&(m.provincia==p)&(m.arena=='Legislativa')
        gl=g[g.porcentaje_Legislativa.notna()|g.escanos.gt(0)]
        m.loc[mask,'gallagher']=math.sqrt(((gl.porcentaje_Legislativa.fillna(0)-gl.escanos_pct.fillna(0))**2).sum()/2)
        m.loc[mask,'voto_sin_escanos_pct']=gl.loc[gl.escanos.eq(0),'porcentaje_Legislativa'].sum()
        ss=gl.escanos.fillna(0);m.loc[mask,'nep_escanos']=1/((ss/ss.sum())**2).sum() if ss.sum() else np.nan
        m.loc[mask,'magnitud']=g.magnitud.max()
        m.loc[mask,'mujeres_electas_pct']=100*g.mujeres_electas.sum()/g.electas_registro.sum() if g.electas_registro.sum() else np.nan
        m.loc[mask,'mujeres_candidatas_pct']=100*g.mujeres_propietarias.sum()/g.propietarias.sum() if g.propietarias.sum() else np.nan
        # Distancia entre distribuciones sobre la unión de partidos; ausencia de oferta=0.
        dual=g[g.porcentaje_Presidencial.notna()|g.porcentaje_Legislativa.notna()]
        m.loc[mask,'divergencia_arenas']=abs(dual.porcentaje_Presidencial.fillna(0)-dual.porcentaje_Legislativa.fillna(0)).sum()/2
    # Perfil longitudinal nominal: no presume continuidad jurídica del nombre.
    return m,panel

def complementary_analysis(v,m,panel,c,e,f,events):
    """Tablas adicionales: asociaciones descriptivas y antecedentes observados."""
    out={}
    histories=[]
    for pt,g in panel[panel.provincia=='Nacional'].groupby('partido'):
        observed=g[g.candidaturas.gt(0)|g.votos_Legislativa.gt(0)|g.votos_Presidencial.gt(0)]
        ys=sorted(observed.anio.unique());first_seat=g[g.escanos.gt(0)].anio.min()
        histories.append({'partido':pt,'primera_aparicion':int(min(ys)),'ultima_aparicion':int(max(ys)),'elecciones_observadas':len(ys),'primera_diputacion':int(first_seat) if pd.notna(first_seat) else None,'escanos_acumulados':g.escanos.sum(),'denominacion':'Etiqueta nominal; continuidad jurídica no establecida'})
    out['historia_partidos']=pd.DataFrame(histories)
    geographic=[]
    firsts={r['partido']:r['primera_aparicion'] for r in histories}
    for (y,pt),g in v[(v.arena=='Legislativa')&(v.provincia!='Nacional')].groupby(['anio','partido']):
        points=g.set_index('provincia').porcentaje.reindex(PROVINCES).fillna(0)
        candg=c[(c.anio==y)&(c.partido==pt)]
        presg=f[(f.anio==y)&(f.partido==pt)]
        geographic.append({'anio':y,'partido':pt,'provincias_con_votos':int(g[g.votos.gt(0)].provincia.nunique()),'provincias_con_candidaturas':int(candg.provincia.nunique()),'candidaturas':len(candg),'propietarias':int(candg.tipo.eq('Propietaria').sum()),'oferta_presidencial':bool(len(presg)),'nombre_nuevo_en_corpus':firsts.get(pt)==y,'dispersion_apoyo_pp':float(points.std(ddof=0)),'maximo_provincial_pct':float(points.max()),'provincia_maximo':points.idxmax(),'votos':float(g.votos.sum())})
    out['alcance_partidos']=pd.DataFrame(geographic)
    gender=[]
    for scope in ['Nacional']+PROVINCES:
        for y in YEARS:
            a=c[c.anio==y];b=e[e.anio==y]
            if scope!='Nacional':a=a[a.provincia==scope];b=b[b.provincia==scope]
            own=a[a.tipo=='Propietaria'];su=a[a.tipo=='Suplencia'];heads=own[own.orden==1]
            ev=events[(events.anio==y)&(events.cargo=='Diputación')]
            if scope!='Nacional':ev=ev[ev.provincia==scope]
            ids=ev.drop_duplicates('id_persona')
            row={'anio':y,'provincia':scope,'propietarias':len(own),'suplencias':len(su),'electas':len(b),'mujeres_propietarias':int(own.sexo.eq('Mujer').sum()),'mujeres_suplentes':int(su.sexo.eq('Mujer').sum()),'mujeres_electas':int(b.sexo.eq('Mujer').sum()),'encabezamientos':len(heads),'encabezamientos_mujeres':int(heads.sexo.eq('Mujer').sum()),'personas_nominales':len(ids),'personas_con_antecedente':int(ids.experiencia_previa.sum()),'nuevas_en_corpus':int((~ids.experiencia_previa).sum())}
            for key,num,den in [('mujeres_propietarias_pct','mujeres_propietarias','propietarias'),('mujeres_suplentes_pct','mujeres_suplentes','suplencias'),('mujeres_electas_pct','mujeres_electas','electas'),('mujeres_encabezamientos_pct','encabezamientos_mujeres','encabezamientos'),('experiencia_pct','personas_con_antecedente','personas_nominales')]:row[key]=ratio(row[num],row[den])
            gender.append(row)
    out['genero_renovacion']=pd.DataFrame(gender)
    govt=[]
    presidents=f[(f.cargo=='Presidencia')&f.electa]
    national=panel[panel.provincia=='Nacional']
    for r in presidents.itertuples():
        current=national[(national.anio==r.anio)&(national.partido==r.partido)]
        if current.empty:continue
        p=current.iloc[0];idx=YEARS.index(r.anio);nextyear=YEARS[idx+1] if idx+1<len(YEARS) else None
        nxt=national[(national.anio==nextyear)&(national.partido==r.partido)]
        ev=events[(events.anio==r.anio)&(events.nombre_key==r.nombre_key)]
        prior=events[(events.anio<r.anio)&(events.nombre_key==r.nombre_key)]
        govt.append({'anio':r.anio,'nombre':r.nombre,'partido':r.partido,'escanos':p.escanos,'escanos_pct':p.escanos_pct,'magnitud':p.magnitud,'mayoria_inicial':bool(p.escanos>p.magnitud/2),'antecedente_candidatura':bool(len(prior)),'antecedente_legislativo_electo':bool(len(prior[(prior.cargo=='Diputación')&prior.electa])),'eleccion_siguiente':nextyear,'cambio_voto_legislativo_pp':float(nxt.iloc[0].porcentaje_Legislativa-p.porcentaje_Legislativa) if len(nxt) else None,'cambio_escanos':float(nxt.iloc[0].escanos-p.escanos) if len(nxt) else None})
    out['presidencia_asamblea']=pd.DataFrame(govt)
    correlations=[]
    cols=['abstencion_pct','invalidos_pct','nep','concentracion_2','margen','gallagher','mujeres_electas_pct','mujeres_candidatas_pct','divergencia_arenas','volatilidad_nominal']
    for scope in ['Nacional','Siete provincias']:
        pop=m[(m.arena=='Legislativa')&(m.ronda==1)&(m.provincia.eq('Nacional') if scope=='Nacional' else m.provincia.ne('Nacional'))]
        for i,x in enumerate(cols):
            for y in cols[i+1:]:
                complete=pop[[x,y]].dropna()
                if len(complete)>=3 and complete[x].nunique()>1 and complete[y].nunique()>1:
                    correlations.append({'ambito':scope,'variable_x':x,'variable_y':y,'observaciones':len(complete),'pearson':complete[x].corr(complete[y]),'spearman':complete[x].rank().corr(complete[y].rank()),'interpretacion':'Exploratoria; dependencia temporal y territorial; sin inferencia causal ni selección por significancia'})
    out['asociaciones']=pd.DataFrame(correlations)
    return out

def build_dataset(root,output):
    output=Path(output);output.mkdir(parents=True,exist_ok=True)
    raw,manifest=load_sources(root);v,t,s,issues=extract_elections(root)
    c,e,f,events,people,moves,placeholders=prepare_people(raw,issues)
    m,panel=calculate_metrics(v,t,s,c,e,events)
    d=raw['diseno'].rename(columns={'Año':'anio','Provincia':'provincia'})
    ca=raw['cancelados'].copy()
    for col in ca.columns:
        if 'Fecha' in str(col):ca[col]=pd.to_datetime(ca[col],errors='coerce').dt.strftime('%Y-%m-%d')
    for y in YEARS:
        n=int((e.anio==y).sum());expected=45 if y in (1953,1958) else 57
        if n!=expected:issues.append({'base':'electas','anio':y,'provincia':'Nacional','detalle':'Total electas frente a magnitud histórica','diferencia':n-expected})
    source_counts={k:len(x) for k,x in raw.items()}
    data={'meta':{'autor':'Juan Ignacio Garbanzo Fallas','perfil':'Politólogo · Economista · Ciencia de Datos Política','proyecto':'Práctica Profesional · IFED / Tribunal Supremo de Elecciones','titulo':'Atlas electoral de Costa Rica','periodo':'1953–2026','version':'2.0','fecha':datetime.now().isoformat(timespec='seconds'),'fuentes':manifest,'filas_fuentes':source_counts,'cobertura':'Siete provincias. Resultados presidenciales según cobertura de las bases aportadas; no se agrega voto exterior.','identidad':'Personas vinculadas por nombre completo normalizado. Trayectorias exploratorias; homónimos no descartados de forma exhaustiva. Partidos comparados por etiqueta nominal, no por continuidad jurídica.','filas_sin_nombre':len(placeholders),'años':YEARS,'provincias':PROVINCES},'votos':records(v),'totales':records(t),'escanos':records(s),'metricas':records(m),'panel':records(panel),'candidaturas':records(c.drop(columns=['nombre_key'],errors='ignore')),'electas':records(e.drop(columns=['nombre_key'],errors='ignore')),'formulas':records(f.drop(columns=['nombre_key'],errors='ignore')),'eventos':records(events.drop(columns=['nombre_key'])),'personas':records(people),'movimientos':records(moves),'papeletas':records(d),'cancelados':records(ca),'integridad':issues}
    from atlas_outcomes import build_outcomes
    data.update({key:records(df) for key,df in build_outcomes(events,e,f).items()})
    data['coincidencias_registrales']=records(raw['coincidencias_registrales'])
    data.update({key:records(table) for key,table in complementary_analysis(v,m,panel,c,e,f,events).items()})
    from atlas_enhancements import enrich
    data=enrich(data)
    from atlas_geography import enrich_geography
    data=enrich_geography(data)
    for key,rows in data.items():
        if isinstance(rows,list):pd.DataFrame(rows).to_csv(output/f'{key}.csv',index=False,encoding='utf-8-sig')
    (output/'datos_integrados.json').write_text(json.dumps(data,ensure_ascii=False,allow_nan=False),encoding='utf-8')
    (output/'manifest_fuentes.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
    return data

def render_app(data,template,output):
    from plotly.offline import get_plotlyjs
    text=Path(template).read_text(encoding='utf-8')
    extensions='\n'.join((Path(template).parent/name).read_text(encoding='utf-8')
                         for name in ('atlas_v3.js','atlas_final.js','atlas_v12.js','atlas_v14.js'))
    text=text.replace('__V3_EXTENSIONS__',extensions)
    payload=json.dumps(data,ensure_ascii=False,allow_nan=False).replace('</','<\\/')
    text=text.replace('__PLOTLY_JS__',get_plotlyjs()).replace('__ATLAS_DATA__',payload)
    required = ['brechas_participacion', 'rankings_provinciales', 'concordancia_gobierno', 'representacion_territorial', 'geografia', 'resultados_candidaturas', 'resultados_personas', 'transiciones_resultados', 'puentes_cargos', 'antecedentes_cargos_electos', 'camara_electa', 'ejecutivo_electo']
    missing = [key for key in required if key not in data]
    if missing:
        raise ValueError("Faltan tablas de la aplicación: " + ", ".join(missing) + ". Ejecute Entorno e Integración de este paquete.")
    unresolved = [token for token in ('__V3_EXTENSIONS__', '__PLOTLY_JS__', '__ATLAS_DATA__') if token in text]
    if unresolved:
        raise ValueError("HTML incompleto: " + ", ".join(unresolved))
    Path(output).parent.mkdir(parents=True,exist_ok=True)
    Path(output).write_text(text,encoding='utf-8')
    return Path(output)
